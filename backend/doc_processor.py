import os
import uuid
import pymupdf as fitz  # PyMuPDF for PDF
import docx
from pptx import Presentation

from langchain.text_splitter import RecursiveCharacterTextSplitter
from qdrant_client import QdrantClient
from qdrant_client.http import models
from sentence_transformers import SentenceTransformer


# ---------- Qdrant Setup ----------
QDRANT_API_KEY = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJhY2Nlc3MiOiJtIn0.5hqNxJ5Lr9KCqmgjCIjA4THBoPw1n_UYaJ3AQTJWNRc"
QDRANT_URL = "https://0896e0e6-4b50-438a-82db-09ba2ce1b1a8.us-east-1-1.aws.cloud.qdrant.io"

client = QdrantClient(url=QDRANT_URL, api_key=QDRANT_API_KEY, timeout=60)
embedding_model = SentenceTransformer("BAAI/bge-small-en-v1.5")


# ---------- Text Extraction ----------
def extract_pdf_text(path: str):
    doc = fitz.open(path)
    return [page.get_text("text") for page in doc if page.get_text("text").strip()]

def extract_docx_text(path: str):
    doc = docx.Document(path)
    return [para.text for para in doc.paragraphs if para.text.strip()]

def extract_pptx_text(path: str):
    prs = Presentation(path)
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text)
    return texts

def extract_text(path: str):
    if path.endswith(".pdf"):
        return extract_pdf_text(path)
    elif path.endswith(".docx"):
        return extract_docx_text(path)
    elif path.endswith(".pptx"):
        return extract_pptx_text(path)
    else:
        raise ValueError("Unsupported file type")


# ---------- Chunking ----------
def chunk_text(pages, chunk_size=800, chunk_overlap=150):
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )

    all_chunks = []
    for page_num, text in enumerate(pages):
        if not text.strip():
            continue

        chunks = splitter.split_text(text)
        for idx, chunk in enumerate(chunks):
            if len(chunk.strip()) < 20:
                continue
            all_chunks.append(
                (
                    chunk,
                    {
                        "page": page_num + 1,
                        "chunk_index": idx,
                        "chunk_length": len(chunk),
                        "source": f"page-{page_num+1}-chunk-{idx}",
                    },
                )
            )
    return all_chunks


# ---------- Store in Qdrant ----------
def store_chunks_in_qdrant(chunks, collection_name, doc_id, batch_size=50):
    texts = [chunk[0] for chunk in chunks]
    metadatas = [chunk[1] for chunk in chunks]
    embeddings = embedding_model.encode(texts).tolist()

    total_points = 0
    for i in range(0, len(texts), batch_size):
        batch_texts = texts[i:i + batch_size]
        batch_metas = metadatas[i:i + batch_size]
        batch_embeds = embeddings[i:i + batch_size]

        points = []
        for text, meta, vector in zip(batch_texts, batch_metas, batch_embeds):
            meta["doc_id"] = doc_id  # ✅ attach doc_id to each chunk
            points.append(
                models.PointStruct(
                    id=uuid.uuid4().int >> 64,
                    vector=vector,
                    payload={"text": text, "doc_id": doc_id, "metadata": meta},
                )
            )

        client.upsert(collection_name=collection_name, points=points)
        total_points += len(points)

    return total_points


# ---------- End-to-End ----------
def process_file_to_qdrant(file_path, collection_name="notes_collection_cloud", doc_id=None):
    if doc_id is None:
        doc_id = str(uuid.uuid4())

    pages = extract_text(file_path)
    chunks = chunk_text(pages)
    stored = store_chunks_in_qdrant(chunks, collection_name, doc_id)
    print(f"✅ {stored} chunks stored for doc_id={doc_id}")

    # ✅ Return both doc_id and stored count
    return {"doc_id": doc_id, "chunks_stored": stored}
